"""将会话记录导出为 Word / PDF。"""

from __future__ import annotations

import io
import re
from datetime import datetime
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt, RGBColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from app.models import Message, Session

# 角色中文名
_ROLE_LABEL = {
    "user": "用户",
    "assistant": "助手",
    "system": "系统",
}


def _role_label(role: str) -> str:
    return _ROLE_LABEL.get(role, role)


def _format_time(value: datetime | None) -> str:
    if not value:
        return ""
    # 统一显示为本地可读格式（去掉微秒）
    if value.tzinfo is not None:
        value = value.astimezone()
    return value.strftime("%Y-%m-%d %H:%M:%S")


def sanitize_filename(title: str, fallback: str = "对话记录") -> str:
    """生成适合作为下载文件名的安全字符串。"""
    text = (title or "").strip() or fallback
    text = re.sub(r'[\\/:*?"<>|\s]+', "_", text)
    text = text.strip("._") or fallback
    return text[:80]


def build_docx(session: Session, messages: list[Message]) -> bytes:
    """生成 Word (.docx) 二进制内容。"""
    doc = Document()

    title = session.title or "新对话"
    heading = doc.add_heading(title, level=0)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x0D, 0x0D, 0x0D)

    meta = doc.add_paragraph()
    meta_run = meta.add_run(
        f"模型：{session.model}　　导出时间：{_format_time(datetime.now().astimezone())}"
    )
    meta_run.font.size = Pt(10)
    meta_run.font.color.rgb = RGBColor(0x6E, 0x6E, 0x80)

    doc.add_paragraph("")

    if not messages:
        doc.add_paragraph("（暂无消息）")
    else:
        for msg in messages:
            role_p = doc.add_paragraph()
            role_run = role_p.add_run(
                f"{_role_label(msg.role)}　{_format_time(msg.created_at)}"
            )
            role_run.bold = True
            role_run.font.size = Pt(11)
            if msg.role == "user":
                role_run.font.color.rgb = RGBColor(0x54, 0x36, 0xDA)
            elif msg.role == "assistant":
                role_run.font.color.rgb = RGBColor(0x10, 0xA3, 0x7F)

            body = doc.add_paragraph()
            body_run = body.add_run(msg.content or "")
            body_run.font.size = Pt(11)
            body.paragraph_format.space_after = Pt(12)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _pdf_escape_text(text: str) -> str:
    """转义后供 ReportLab Paragraph 使用，保留换行。"""
    return escape(text or "").replace("\n", "<br/>")


def build_pdf(session: Session, messages: list[Message]) -> bytes:
    """生成 PDF 二进制内容（中文字体使用内置 CID 字体）。"""
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title=session.title or "对话记录",
    )

    title_style = ParagraphStyle(
        "TitleCN",
        fontName="STSong-Light",
        fontSize=18,
        leading=24,
        spaceAfter=8,
        textColor="#0d0d0d",
    )
    meta_style = ParagraphStyle(
        "MetaCN",
        fontName="STSong-Light",
        fontSize=9,
        leading=14,
        textColor="#6e6e80",
        spaceAfter=16,
    )
    role_style = ParagraphStyle(
        "RoleCN",
        fontName="STSong-Light",
        fontSize=11,
        leading=16,
        textColor="#343541",
        spaceBefore=10,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "BodyCN",
        fontName="STSong-Light",
        fontSize=10.5,
        leading=16,
        textColor="#0d0d0d",
        spaceAfter=8,
    )

    story = [
        Paragraph(_pdf_escape_text(session.title or "新对话"), title_style),
        Paragraph(
            _pdf_escape_text(
                f"模型：{session.model}　　导出时间：{_format_time(datetime.now().astimezone())}"
            ),
            meta_style,
        ),
    ]

    if not messages:
        story.append(Paragraph("（暂无消息）", body_style))
    else:
        for msg in messages:
            story.append(
                Paragraph(
                    _pdf_escape_text(
                        f"{_role_label(msg.role)}　{_format_time(msg.created_at)}"
                    ),
                    role_style,
                )
            )
            story.append(Paragraph(_pdf_escape_text(msg.content or ""), body_style))
            story.append(Spacer(1, 4))

    doc.build(story)
    return buffer.getvalue()
